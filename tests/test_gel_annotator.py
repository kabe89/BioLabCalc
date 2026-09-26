import unittest
import os
from PIL import Image
import openpyxl
from biolabcalc.gel_annotator import GelAnnotator

class TestGelAnnotator(unittest.TestCase):
    def setUp(self):
        self.output_png = "/tmp/test_annotator_out.png"
        self.output_xlsx = "/tmp/test_annotator_out.xlsx"
        self.test_raw_img = "/tmp/test_raw_gel.png"

        # Create dummy raw gel image
        img = Image.new("L", (400, 300), color=50)
        img.save(self.test_raw_img)

    def tearDown(self):
        for p in (self.output_png, self.output_xlsx, self.test_raw_img):
            if os.path.exists(p):
                try:
                    os.remove(p)
                except Exception:
                    pass

    def test_synthetic_annotator_render(self):
        ann = GelAnnotator.synthetic(
            num_lanes=5,
            title="Synthetic PCR Test Gel",
            gel_info="1.0% Agarose",
        )
        ann.set_lanes(["Ladder", "Ctrl", "Sample 1", "Sample 2", "Sample 3"])
        ann.add_ladder(lane_index=1, ladder_type="1kb_dna", side="left")
        ann.add_band(lane_index=3, size_bp_or_kda=750, label="750 bp Product", color="red")
        ann.add_band(lane_index=4, size_bp_or_kda=1500, label="1.5 kb Product", color="blue")

        res_path = ann.render(self.output_png, dpi=150)
        self.assertTrue(os.path.exists(res_path))
        self.assertGreater(os.path.getsize(res_path), 1000)

    def test_from_image_annotator_render(self):
        ann = GelAnnotator.from_image(
            image_path=self.test_raw_img,
            title="Real Gel Photo Analysis",
            invert_colors=True,
        )
        ann.set_lanes(["Marker", "Lysate", "Eluate"])
        ann.add_ladder(lane_index=1, ladder_type="protein_broad_range", side="left")
        ann.add_band(lane_index=3, size_bp_or_kda=50, label="50 kDa Target Protein")

        res_path = ann.render(self.output_png, dpi=150)
        self.assertTrue(os.path.exists(res_path))
        self.assertGreater(os.path.getsize(res_path), 1000)

    def test_embed_in_excel(self):
        ann = GelAnnotator.synthetic(num_lanes=4)
        ann.set_lanes(["Ladder", "A", "B", "C"])
        ann.add_ladder(1, "100bp_dna")
        ann.add_band(2, 500, "500 bp Band")

        res_xl = ann.embed_in_excel(self.output_xlsx, sheet_name="Annotated_Gel")
        self.assertTrue(os.path.exists(res_xl))
        wb = openpyxl.load_workbook(res_xl)
        self.assertIn("Annotated_Gel", wb.sheetnames)
        ws = wb["Annotated_Gel"]
        # Check lane headers written in table
        self.assertEqual(ws.cell(row=6, column=2).value, "Ladder")
        self.assertEqual(ws.cell(row=7, column=2).value, "A")
        # Check image was added
        self.assertGreater(len(ws._images), 0)

if __name__ == '__main__':
    unittest.main()
