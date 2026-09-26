import unittest
import os
import openpyxl
from biolabcalc.excel_extension import generate_lab_notebook_template, export_analysis_to_excel, batch_process_excel
from biolabcalc.transcription import calculate_ivt_yield
from biolabcalc.pcr import calculate_pcr_kinetics

class TestExcelExtension(unittest.TestCase):
    def test_generate_template(self):
        out_path = "/tmp/test_gen_template.xlsx"
        res_path = generate_lab_notebook_template(out_path)
        self.assertTrue(os.path.exists(res_path))
        wb = openpyxl.load_workbook(res_path)
        self.assertIn("IVT_Stoichiometry", wb.sheetnames)
        self.assertIn("PCR_Optimization", wb.sheetnames)
        self.assertIn("Protein_Quantification", wb.sheetnames)
        self.assertIn("Primer_Design_Log", wb.sheetnames)

    def test_export_analysis(self):
        out_path = "/tmp/test_export_analysis.xlsx"
        ivt = calculate_ivt_yield("GGGAAUGCAUGCAUGC", measured_yield_ug=25.0)
        pcr = calculate_pcr_kinetics(amplicon_len_bp=400, template_ng=5.0)
        export_analysis_to_excel(out_path, ivt_results=[ivt], pcr_results=[pcr])
        self.assertTrue(os.path.exists(out_path))
        wb = openpyxl.load_workbook(out_path)
        self.assertIn("IVT_Results", wb.sheetnames)
        self.assertIn("PCR_Results", wb.sheetnames)


    def test_batch_process_excel(self):
        in_path = "/tmp/test_unit_batch_in.xlsx"
        out_path = "/tmp/test_unit_batch_out.xlsx"
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append(["ID", "Sequence", "Type"])
        ws.append(["Oligo_1", "ATGCCGTCCAGGCTGCTGGTC", "DNA"])
        ws.append(["RNA_1", "GGGAGACCCAAGCUGGCCUGUGUGUACAGCCGAAUG", "RNA"])
        wb.save(in_path)

        res = batch_process_excel(in_path, out_path)
        self.assertTrue(os.path.exists(res))
        out_wb = openpyxl.load_workbook(res)
        self.assertIn("Batch_Analysis_Results", out_wb.sheetnames)
        ws_out = out_wb["Batch_Analysis_Results"]
        self.assertEqual(ws_out.cell(row=4, column=2).value, "Oligo_1")
        self.assertGreater(ws_out.cell(row=4, column=6).value, 5000)

if __name__ == '__main__':
    unittest.main()
