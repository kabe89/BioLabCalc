"""Laboratory report data recorder, statistical aggregator, and multi-format exporter (Markdown, Excel, CSV, JSON).

Provides a structured, user-friendly system for bench researchers to record experimental measurements,
track replicates, compute statistical metrics (Mean, SD, %CV), evaluate QC thresholds, and export
publication-ready reports for Electronic Lab Notebooks (ELN) and lab archives.
"""

from __future__ import annotations
import os
import csv
import json
import math
from datetime import datetime
from dataclasses import dataclass, field, asdict
from typing import Dict, List, Optional, Any, Union


@dataclass
class MeasurementRecord:
    sample_id: str
    parameter: str
    value: float
    unit: str
    target_or_expected: Optional[float] = None
    tolerance_pct: Optional[float] = None
    status: str = "PASS"  # PASS, FLAG, FAIL, INFO
    timestamp: str = field(default_factory=lambda: datetime.now().strftime("%Y-%m-%d %H:%M:%S"))
    notes: str = ""

    @property
    def percent_deviation(self) -> Optional[float]:
        if self.target_or_expected is not None and self.target_or_expected != 0:
            return round(((self.value - self.target_or_expected) / self.target_or_expected) * 100.0, 2)
        return None


@dataclass
class ParameterStatistics:
    sample_id: str
    parameter: str
    unit: str
    count: int
    mean: float
    std_dev: float
    cv_percent: float
    min_value: float
    max_value: float


class LabReport:
    """Electronic lab report session manager for recording and exporting bench data."""

    def __init__(
        self,
        title: str = "Laboratory Experiment Report",
        experimenter: str = "Researcher",
        project: str = "General",
        objective: str = "",
        date: Optional[str] = None,
    ):
        self.title = title
        self.experimenter = experimenter
        self.project = project
        self.objective = objective
        self.date = date or datetime.now().strftime("%Y-%m-%d %H:%M")
        self.records: List[MeasurementRecord] = []
        self.conclusions: List[str] = []
        self.metadata: Dict[str, Any] = {}

    def add_record(
        self,
        sample_id: str,
        parameter: str,
        value: float,
        unit: str,
        target: Optional[float] = None,
        tolerance_pct: Optional[float] = None,
        status: Optional[str] = None,
        notes: str = "",
    ) -> MeasurementRecord:
        """Record an experimental measurement with optional QC target and tolerance."""
        val_f = float(value)
        tgt_f = float(target) if target is not None else None

        # Automatic QC status evaluation if target & tolerance provided
        if status is None:
            if tgt_f is not None and tolerance_pct is not None and tgt_f != 0:
                dev = abs((val_f - tgt_f) / tgt_f) * 100.0
                if dev <= tolerance_pct:
                    eval_status = "PASS"
                elif dev <= tolerance_pct * 1.5:
                    eval_status = "FLAG"
                else:
                    eval_status = "FAIL"
            else:
                eval_status = "PASS"
        else:
            eval_status = status.upper()

        rec = MeasurementRecord(
            sample_id=str(sample_id).strip(),
            parameter=str(parameter).strip(),
            value=val_f,
            unit=str(unit).strip(),
            target_or_expected=tgt_f,
            tolerance_pct=tolerance_pct,
            status=eval_status,
            notes=notes.strip(),
        )
        self.records.append(rec)
        return rec

    def record_nanodrop(
        self,
        sample_id: str,
        purity_or_solution_result: Any,
        notes: str = "",
    ) -> List[MeasurementRecord]:
        """Convenience method to log NanoDrop purity or standardized solution results."""
        added = []
        if hasattr(purity_or_solution_result, "a260"):
            res = purity_or_solution_result
            status = "PASS" if getattr(res, "purity_status", "Pure") == "Pure" else "FLAG"
            added.append(self.add_record(sample_id, "A260", res.a260, "AU", notes=notes))
            added.append(self.add_record(sample_id, "A260/A280", res.a260_a280_ratio, "ratio", target=2.0 if "RNA" in getattr(res, "sample_type", "RNA") else 1.8, tolerance_pct=15.0, status=status))
            added.append(self.add_record(sample_id, "A260/A230", res.a260_a230_ratio, "ratio", target=2.1, tolerance_pct=20.0))
            if hasattr(res, "corrected_concentration_ng_ul"):
                added.append(self.add_record(sample_id, "Corrected Concentration", res.corrected_concentration_ng_ul, "ng/µL", notes=f"Purity: {getattr(res, 'purity_status', 'Pure')}"))
        elif hasattr(purity_or_solution_result, "actual_concentration_ng_ul"):
            res = purity_or_solution_result
            added.append(self.add_record(sample_id, "Concentration", res.actual_concentration_ng_ul, "ng/µL", notes=notes))
            added.append(self.add_record(sample_id, "Expected A260", res.expected_nanodrop_a260_sequence_specific, "AU"))
            added.append(self.add_record(sample_id, "Molecular Weight", res.molecular_weight, "Da"))
        return added

    def record_ivt(
        self,
        sample_id: str,
        ivt_result: Any,
        notes: str = "",
    ) -> List[MeasurementRecord]:
        """Convenience method to log IVT yield, efficiency, and reaction stoichiometry."""
        added = []
        res = ivt_result
        added.append(self.add_record(sample_id, "RNA Yield", res.rna_yield_ug, "µg", target=res.theoretical_max_yield_ug, notes=notes))
        added.append(self.add_record(sample_id, "Incorporation Efficiency", res.overall_efficiency_percent, "%", target=100.0))
        added.append(self.add_record(sample_id, "Pyrophosphate (PPi)", res.pyrophosphate_released_ug, "µg"))
        added.append(self.add_record(sample_id, "Free Initial Mg2+", res.free_mg_initial_mm, "mM", notes=f"Initiation: {res.initiation_efficiency_rating}"))
        if hasattr(res, "protons_released_nmol"):
            added.append(self.add_record(sample_id, "Protons Released", res.protons_released_nmol, "nmol"))
        return added

    def record_pcr(
        self,
        sample_id: str,
        pcr_result: Any,
        notes: str = "",
    ) -> List[MeasurementRecord]:
        """Convenience method to log PCR kinetics and amplicon yield."""
        added = []
        res = pcr_result
        added.append(self.add_record(sample_id, "Amplicon Length", float(res.amplicon_length_bp), "bp", notes=notes))
        added.append(self.add_record(sample_id, "Amplicon Yield", res.amplicon_yield_ng, "ng", target=res.theoretical_max_amplicon_ug * 1000.0))
        added.append(self.add_record(sample_id, "PCR Cycles", float(res.cycles_run), "cycles"))
        if res.dntp_exhaustion_cycle:
            added.append(self.add_record(sample_id, "Plateau Cycle", float(res.dntp_exhaustion_cycle), "cycle"))
        return added

    def record_western(
        self,
        sample_id: str,
        band_mw_kda: float,
        intensity_au: float,
        target_kda: Optional[float] = None,
        notes: str = "",
    ) -> List[MeasurementRecord]:
        """Convenience method to log Western blot band sizing and densitometry intensity."""
        added = []
        added.append(self.add_record(sample_id, "Observed MW", float(band_mw_kda), "kDa", target=target_kda, tolerance_pct=10.0 if target_kda else None, notes=notes))
        added.append(self.add_record(sample_id, "Band Densitometry", float(intensity_au), "AU"))
        return added

    def add_conclusion(self, statement: str) -> None:
        """Add a summary takeaway or conclusion to the report."""
        if statement.strip():
            self.conclusions.append(statement.strip())

    def calculate_statistics(
        self,
        sample_id: Optional[str] = None,
        parameter: Optional[str] = None,
    ) -> List[ParameterStatistics]:
        """Compute summary statistics (Mean, SD, %CV, Min, Max) grouped by sample and parameter."""
        grouped: Dict[Tuple[str, str, str], List[float]] = {}
        for r in self.records:
            if sample_id and r.sample_id != sample_id:
                continue
            if parameter and r.parameter != parameter:
                continue
            key = (r.sample_id, r.parameter, r.unit)
            if key not in grouped:
                grouped[key] = []
            grouped[key].append(r.value)

        stats_list = []
        for (s_id, param, unit), vals in grouped.items():
            n = len(vals)
            mean_val = sum(vals) / n
            if n > 1:
                variance = sum((x - mean_val) ** 2 for x in vals) / (n - 1)
                std_dev = math.sqrt(variance)
                cv = (std_dev / mean_val) * 100.0 if mean_val != 0 else 0.0
            else:
                std_dev = 0.0
                cv = 0.0

            stats_list.append(ParameterStatistics(
                sample_id=s_id,
                parameter=param,
                unit=unit,
                count=n,
                mean=round(mean_val, 3),
                std_dev=round(std_dev, 3),
                cv_percent=round(cv, 2),
                min_value=round(min(vals), 3),
                max_value=round(max(vals), 3),
            ))

        return stats_list

    def to_markdown(self, filepath: Optional[str] = None) -> str:
        """Generate a complete, publication-grade Markdown laboratory report."""
        lines = [
            f"# 🧪 {self.title}",
            "",
            f"**Researcher:** {self.experimenter}  |  **Project:** {self.project}  |  **Date:** {self.date}",
        ]
        if self.objective:
            lines.extend([
                "",
                "## 🎯 Objective & Experimental Scope",
                self.objective,
            ])

        lines.extend([
            "",
            "## 📊 Recorded Experimental Data & Measurements",
            "| Sample ID | Assay / Parameter | Value | Unit | Target | Dev (%) | QC Status | Notes |",
            "| :--- | :--- | :--- | :--- | :--- | :--- | :---: | :--- |",
        ])

        for r in self.records:
            tgt_str = f"{r.target_or_expected:g}" if r.target_or_expected is not None else "-"
            dev_str = f"{r.percent_deviation:+.1f}%" if r.percent_deviation is not None else "-"
            status_badge = f"**{r.status}**" if r.status != "PASS" else "PASS"
            lines.append(
                f"| `{r.sample_id}` | {r.parameter} | **{r.value:g}** | {r.unit} | {tgt_str} | {dev_str} | {status_badge} | {r.notes} |"
            )

        # Statistical Summary Table (for parameters with replicates)
        stats = self.calculate_statistics()
        reps = [s for s in stats if s.count > 1]
        if reps:
            lines.extend([
                "",
                "## 📈 Replicate Statistics (Mean, SD, %CV)",
                "| Sample ID | Parameter | N | Mean ± SD | %CV | Range [Min, Max] |",
                "| :--- | :--- | :---: | :--- | :---: | :--- |",
            ])
            for s in reps:
                lines.append(
                    f"| `{s.sample_id}` | {s.parameter} | {s.count} | **{s.mean:g} ± {s.std_dev:g}** {s.unit} | {s.cv_percent:.1f}% | [{s.min_value:g}, {s.max_value:g}] |"
                )

        if self.conclusions:
            lines.extend([
                "",
                "## 📝 Key Conclusions & Observations",
            ])
            for idx, c in enumerate(self.conclusions, 1):
                lines.append(f"{idx}. {c}")

        lines.append("")
        md_text = "\n".join(lines)

        if filepath:
            with open(filepath, "w", encoding="utf-8") as f:
                f.write(md_text)

        return md_text

    def to_excel(self, filepath: str) -> str:
        """Export the lab report to a styled Microsoft Excel workbook with formulas and QC formatting."""
        import openpyxl
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
        from openpyxl.utils import get_column_letter

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Lab_Report"
        ws.views.sheetView[0].showGridLines = True

        # Styles
        navy_header = PatternFill(start_color="1F497D", end_color="1F497D", fill_type="solid")
        teal_sub = PatternFill(start_color="008080", end_color="008080", fill_type="solid")
        light_gray = PatternFill(start_color="F2F2F2", end_color="F2F2F2", fill_type="solid")
        pass_fill = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")  # soft green
        flag_fill = PatternFill(start_color="FFEB9C", end_color="FFEB9C", fill_type="solid")  # soft yellow
        fail_fill = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")  # soft red

        white_bold = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        title_font = Font(name="Calibri", size=16, bold=True, color="1F497D")
        sub_font = Font(name="Calibri", size=10, italic=True, color="595959")
        bold_font = Font(name="Calibri", size=11, bold=True)
        regular_font = Font(name="Calibri", size=11)

        thin_border = Border(
            left=Side(style="thin", color="D9D9D9"),
            right=Side(style="thin", color="D9D9D9"),
            top=Side(style="thin", color="D9D9D9"),
            bottom=Side(style="thin", color="D9D9D9"),
        )

        # 1. Header Banner
        ws.merge_cells("A1:H1")
        ws["A1"] = f"🔬 {self.title}"
        ws["A1"].font = title_font
        ws["A1"].alignment = Alignment(vertical="center")

        ws.merge_cells("A2:H2")
        ws["A2"] = f"Researcher: {self.experimenter}   |   Project: {self.project}   |   Date: {self.date}"
        ws["A2"].font = sub_font

        if self.objective:
            ws.merge_cells("A3:H3")
            ws["A3"] = f"Objective: {self.objective}"
            ws["A3"].font = regular_font
            start_row = 5
        else:
            start_row = 4

        # 2. Data Table
        headers = ["Sample ID", "Parameter / Assay", "Measured Value", "Unit", "Target Value", "Deviation (%)", "QC Status", "Bench Notes"]
        for col_idx, h in enumerate(headers, 1):
            cell = ws.cell(row=start_row, column=col_idx, value=h)
            cell.fill = navy_header
            cell.font = white_bold
            cell.alignment = Alignment(horizontal="center", vertical="center")

        curr_row = start_row + 1
        for r in self.records:
            c1 = ws.cell(row=curr_row, column=1, value=r.sample_id)
            c2 = ws.cell(row=curr_row, column=2, value=r.parameter)
            c3 = ws.cell(row=curr_row, column=3, value=r.value)
            c4 = ws.cell(row=curr_row, column=4, value=r.unit)
            c5 = ws.cell(row=curr_row, column=5, value=r.target_or_expected if r.target_or_expected is not None else "")
            
            # Excel formula for deviation: =IF(E{row}<>"", (C{row}-E{row})/E{row}, "")
            c6 = ws.cell(row=curr_row, column=6)
            if r.target_or_expected is not None:
                c6.value = f"=(C{curr_row}-E{curr_row})/E{curr_row}"
                c6.number_format = "+0.0%;-0.0%;0.0%"
            else:
                c6.value = ""

            c7 = ws.cell(row=curr_row, column=7, value=r.status)
            c8 = ws.cell(row=curr_row, column=8, value=r.notes)

            # Alignment and fonts
            for col in range(1, 9):
                cell = ws.cell(row=curr_row, column=col)
                cell.font = regular_font
                cell.border = thin_border
                if col in (3, 5):
                    cell.alignment = Alignment(horizontal="right")
                elif col in (1, 4, 6, 7):
                    cell.alignment = Alignment(horizontal="center")

            # QC Status coloring
            if r.status == "PASS":
                c7.fill = pass_fill
            elif r.status == "FLAG":
                c7.fill = flag_fill
            elif r.status == "FAIL":
                c7.fill = fail_fill

            curr_row += 1

        # 3. Statistics Table
        curr_row += 2
        ws.cell(row=curr_row, column=1, value="📈 SUMMARY REPLICATE STATISTICS").font = Font(name="Calibri", size=13, bold=True, color="008080")
        curr_row += 1

        stat_headers = ["Sample ID", "Parameter", "Unit", "N", "Mean", "Std Dev", "%CV", "Min", "Max"]
        for col_idx, sh in enumerate(stat_headers, 1):
            cell = ws.cell(row=curr_row, column=col_idx, value=sh)
            cell.fill = teal_sub
            cell.font = white_bold
            cell.alignment = Alignment(horizontal="center")

        curr_row += 1
        stats = self.calculate_statistics()
        for s in stats:
            ws.cell(row=curr_row, column=1, value=s.sample_id).alignment = Alignment(horizontal="center")
            ws.cell(row=curr_row, column=2, value=s.parameter)
            ws.cell(row=curr_row, column=3, value=s.unit).alignment = Alignment(horizontal="center")
            ws.cell(row=curr_row, column=4, value=s.count).alignment = Alignment(horizontal="center")
            ws.cell(row=curr_row, column=5, value=s.mean)
            ws.cell(row=curr_row, column=6, value=s.std_dev)
            c_cv = ws.cell(row=curr_row, column=7, value=s.cv_percent / 100.0)
            c_cv.number_format = "0.0%"
            c_cv.alignment = Alignment(horizontal="center")
            ws.cell(row=curr_row, column=8, value=s.min_value)
            ws.cell(row=curr_row, column=9, value=s.max_value)

            for col in range(1, 10):
                cell = ws.cell(row=curr_row, column=col)
                cell.font = regular_font
                cell.border = thin_border
            curr_row += 1

        # 4. Conclusions Block
        if self.conclusions:
            curr_row += 2
            ws.cell(row=curr_row, column=1, value="📝 KEY CONCLUSIONS & OBSERVATIONS").font = Font(name="Calibri", size=13, bold=True, color="1F497D")
            curr_row += 1
            for c in self.conclusions:
                ws.cell(row=curr_row, column=1, value=f"• {c}").font = regular_font
                curr_row += 1

        # Auto-adjust column widths
        for col in ws.columns:
            max_len = 0
            col_letter = get_column_letter(col[0].column)
            for cell in col:
                val = str(cell.value or "")
                if not cell.coordinate in ws.merged_cells:
                    max_len = max(max_len, len(val))
            ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

        wb.save(filepath)
        return filepath

    def to_csv(self, filepath: str) -> str:
        """Export raw measurement records to CSV."""
        with open(filepath, "w", newline="", encoding="utf-8") as f:
            writer = csv.writer(f)
            writer.writerow(["Sample_ID", "Parameter", "Value", "Unit", "Target", "Deviation_Percent", "QC_Status", "Timestamp", "Notes"])
            for r in self.records:
                writer.writerow([
                    r.sample_id,
                    r.parameter,
                    r.value,
                    r.unit,
                    r.target_or_expected if r.target_or_expected is not None else "",
                    r.percent_deviation if r.percent_deviation is not None else "",
                    r.status,
                    r.timestamp,
                    r.notes,
                ])
        return filepath

    def to_json(self, filepath: Optional[str] = None) -> str:
        """Export report and records to JSON string or file."""
        data = {
            "title": self.title,
            "experimenter": self.experimenter,
            "project": self.project,
            "date": self.date,
            "objective": self.objective,
            "records": [asdict(r) for r in self.records],
            "statistics": [asdict(s) for s in self.calculate_statistics()],
            "conclusions": self.conclusions,
        }
        json_str = json.dumps(data, indent=2)
        if filepath:
            with open(filepath, "w", encoding="utf-8") as f:
                f.write(json_str)
        return json_str


# Global default active session for easy interactive / script recording
_DEFAULT_REPORT: LabReport = LabReport()


def get_active_report() -> LabReport:
    """Return the global default active LabReport session."""
    return _DEFAULT_REPORT


def new_lab_report(
    title: str = "Laboratory Experiment Report",
    experimenter: str = "Researcher",
    project: str = "General",
    objective: str = "",
) -> LabReport:
    """Initialize a fresh active LabReport session."""
    global _DEFAULT_REPORT
    _DEFAULT_REPORT = LabReport(
        title=title,
        experimenter=experimenter,
        project=project,
        objective=objective,
    )
    return _DEFAULT_REPORT


def record_measurement(
    sample_id: str,
    parameter: str,
    value: float,
    unit: str,
    target: Optional[float] = None,
    tolerance_pct: Optional[float] = None,
    notes: str = "",
) -> MeasurementRecord:
    """Quick one-line function to record a measurement into the active lab report."""
    return _DEFAULT_REPORT.add_record(
        sample_id=sample_id,
        parameter=parameter,
        value=value,
        unit=unit,
        target=target,
        tolerance_pct=tolerance_pct,
        notes=notes,
    )
